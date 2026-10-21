"""Чтение xlsx со списком заправок.

Ожидаемая структура листа: первая строка — заголовки, далее данные.
Используемые колонки: ``docdate``, ``azscode``, ``trk``, ``cardcode``, ``realdose``.
Остальные колонки (emtcodefrm, contrcode, shiftdate) игнорируются.

Битые строки не выбрасываются: в :attr:`RowData.error` записывается причина,
обработка остальных строк продолжается.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

LOGGER = logging.getLogger(__name__)

#: Колонки, без которых строка не может быть обработана.
REQUIRED_COLUMNS = ("docdate", "azscode", "trk")

#: Поддерживаемые текстовые форматы даты/времени.
DATETIME_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%y %H:%M:%S",
    "%d.%m.%y %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%dT%H:%M",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
)

#: Поправка Excel для числовых дат (1900-я система счёта).
EXCEL_EPOCH = datetime(1899, 12, 30)


class XlsxError(Exception):
    """Ошибка чтения xlsx, текст показывается пользователю."""


@dataclass
class RowData:
    """Одна строка xlsx (одна заправка)."""

    index: int
    azs: Optional[int] = None
    trk: Optional[int] = None
    docdate: Optional[datetime] = None
    cardcode: str = ""
    realdose: Optional[float] = None
    error: str = ""
    raw: Dict[str, Any] = field(default_factory=dict)

    @property
    def is_valid(self) -> bool:
        """Строка пригодна к обработке (нет ошибки разбора)."""
        return not self.error

    @property
    def docdate_text(self) -> str:
        """Дата/время в формате ``ДД.ММ.ГГГГ ЧЧ:ММ:СС`` или «—»."""
        if self.docdate is None:
            return "—"
        return self.docdate.strftime("%d.%m.%Y %H:%M:%S")


def _normalize_header(value: Any) -> str:
    """Привести заголовок колонки к нижнему регистру без пробелов и BOM."""
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").strip().lower()


def _to_int(value: Any) -> Optional[int]:
    """Преобразовать значение ячейки в целое число (``75``, ``75.0``, «ТРК 1»)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(round(value))
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        pass
    digits = "".join(ch for ch in text if ch.isdigit())
    return int(digits) if digits else None


def _to_float(value: Any) -> Optional[float]:
    """Преобразовать значение ячейки в число с плавающей точкой."""
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_datetime(value: Any) -> Optional[datetime]:
    """Преобразовать значение ячейки в ``datetime`` (дата, строка или число Excel)."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return EXCEL_EPOCH + timedelta(days=float(value))
        except (OverflowError, ValueError):
            return None

    text = str(value).strip()
    if not text:
        return None
    for fmt in DATETIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    # Последняя попытка: ISO-формат с секундами/долями секунды.
    try:
        return datetime.fromisoformat(text.replace("Z", ""))
    except ValueError:
        return None


def read_rows(path: Path, sheet_index: int = 0) -> List[RowData]:
    """Прочитать xlsx и вернуть список строк.

    :param path: путь к xlsx
    :param sheet_index: индекс листа (по умолчанию первый)
    :raises XlsxError: если файл не читается или нет обязательных колонок
    :return: список :class:`RowData` (включая строки с ошибками)
    """
    path = Path(path)
    if not path.exists():
        raise XlsxError(f"Файл не найден: {path}")

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl кидает разные типы исключений
        raise XlsxError(f"Не удалось открыть файл {path.name}: {exc}") from exc

    try:
        sheets = workbook.sheetnames
        if not sheets:
            raise XlsxError(f"В файле {path.name} нет ни одного листа.")
        if sheet_index >= len(sheets):
            raise XlsxError(
                f"В файле {path.name} нет листа с индексом {sheet_index}. "
                f"Доступные листы: {', '.join(sheets)}."
            )
        sheet = workbook.worksheets[sheet_index]

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise XlsxError(f"Лист «{sheet.title}» файла {path.name} пуст.")

        columns: Dict[str, int] = {}
        for position, header in enumerate(header_row):
            name = _normalize_header(header)
            if name and name not in columns:
                columns[name] = position

        missing = [name for name in REQUIRED_COLUMNS if name not in columns]
        if missing:
            raise XlsxError(
                f"В файле {path.name} не найдены обязательные колонки: {', '.join(missing)}. "
                f"Найденные заголовки: {', '.join(columns.keys()) or 'нет'}."
            )

        result: List[RowData] = []
        for row_number, values in enumerate(rows_iter, start=2):
            if values is None:
                continue
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            result.append(_build_row(row_number, values, columns))

        LOGGER.info("Прочитано строк из %s: %s", path.name, len(result))
        return result
    finally:
        workbook.close()


def _cell(values: Any, columns: Dict[str, int], name: str) -> Any:
    """Достать значение ячейки по имени колонки."""
    position = columns.get(name)
    if position is None:
        return None
    if position >= len(values):
        return None
    return values[position]


def _build_row(row_number: int, values: Any, columns: Dict[str, int]) -> RowData:
    """Собрать :class:`RowData` из строки листа, отметив ошибки разбора."""
    docdate = _to_datetime(_cell(values, columns, "docdate"))
    azs = _to_int(_cell(values, columns, "azscode"))
    trk = _to_int(_cell(values, columns, "trk"))
    cardcode_value = _cell(values, columns, "cardcode")
    realdose = _to_float(_cell(values, columns, "realdose"))

    errors: List[str] = []
    if docdate is None:
        errors.append("не удалось прочитать docdate")
    if azs is None:
        errors.append("не удалось прочитать azscode")
    if trk is None:
        errors.append("не удалось прочитать trk")

    row = RowData(
        index=row_number,
        azs=azs,
        trk=trk,
        docdate=docdate,
        cardcode=str(cardcode_value).strip() if cardcode_value is not None else "",
        realdose=realdose,
        raw={
            _normalize_header(name): value
            for name, value in zip(columns.keys(), [values[i] for i in columns.values()])
        },
    )
    if errors:
        row.error = "; ".join(errors)
        LOGGER.warning("Строка %s в xlsx пропущена: %s", row_number, row.error)
    return row
