"""Чтение и запись xlsx для вкладки «Ручное копирование».

Формат файла ОТЛИЧАЕТСЯ от xlsx вкладки «Задания»::

    azscode | start             | end               | trk
    40      | 2026-09-17 13:20  | 2026-09-17 13:40  | 2
    40      | 2026-09-17 15:10  | 2026-09-17 15:25  | 1,3

* ``azscode`` — номер АЗС; пусто = текущая АЗС;
* ``start``, ``end`` — дата-время ``YYYY-MM-DD HH:MM`` (при чтении принимаются и
  настоящие даты Excel — они приводятся к тому же виду без секунд);
* ``trk`` — номера ТРК через запятую; пусто = все ТРК из ``azs.ini``.

Модуль не меняет формат файла заданий и не использует его логику.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from xlsx_reader import XlsxError

LOGGER = logging.getLogger(__name__)

#: Формат даты-времени ручного копирования (без секунд).
DATETIME_FORMAT = "%Y-%m-%d %H:%M"

#: Колонки ручного xlsx (первая строка — заголовки).
MANUAL_COLUMNS = ("azscode", "start", "end", "trk")

#: Поправка Excel для числовых дат (1900-я система счёта).
EXCEL_EPOCH = datetime(1899, 12, 30)


@dataclass
class ManualRow:
    """Одна строка ручного задания (в т.ч. прочитанная из xlsx)."""

    azscode: Optional[int]
    start: datetime
    end: datetime
    trks: List[int] = field(default_factory=list)
    error: Optional[str] = None
    status: str = "ожидание"

    @property
    def start_text(self) -> str:
        """Начало интервала в формате ввода."""
        return self.start.strftime(DATETIME_FORMAT)

    @property
    def end_text(self) -> str:
        """Конец интервала в формате ввода."""
        return self.end.strftime(DATETIME_FORMAT)

    @property
    def trks_text(self) -> str:
        """Список ТРК через запятую."""
        return ",".join(str(trk) for trk in self.trks)


def _normalize_header(value: Any) -> str:
    """Привести заголовок колонки к нижнему регистру без пробелов и BOM."""
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").strip().lower()


def parse_dt_text(value: Any) -> Optional[datetime]:
    """Разобрать дату-время в формате ``YYYY-MM-DD HH:MM``.

    Принимает строку, ``datetime``/``date`` из Excel и числовую дату Excel.
    Секунды и микросекунды отбрасываются: ручной интервал задаётся в минутах.

    :return: ``datetime`` или ``None``, если значение не разобрать
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (EXCEL_EPOCH + timedelta(days=float(value))).replace(
                second=0, microsecond=0
            )
        except (OverflowError, ValueError):
            return None

    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, DATETIME_FORMAT)
    except ValueError:
        return None


def _to_int(value: Any) -> Optional[int]:
    """Преобразовать значение ячейки в целое число (``40``, ``40.0``, «40»)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(round(value))
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def _parse_trks(value: Any) -> Tuple[List[int], Optional[str]]:
    """Разобрать ТРК из ячейки: ``"1,3"`` -> ``[1, 3]``.

    :return: кортеж (список ТРК, текст ошибки или ``None``)
    """
    if value is None:
        return [], None
    text = str(value).strip()
    if not text:
        return [], None

    trks: List[int] = []
    for chunk in text.replace(";", ",").split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        try:
            trk = int(chunk)
        except ValueError:
            return trks, f"не удалось прочитать trk: «{chunk}»"
        if trk not in trks:
            trks.append(trk)
    return trks, None


def _open_workbook(path: Path):
    """Открыть xlsx только для чтения или поднять :class:`XlsxError`."""
    try:
        return load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl кидает разные типы исключений
        raise XlsxError(f"Не удалось открыть файл {path.name}: {exc}") from exc


def _read_header(
    workbook: Any, path: Path, sheet_index: int = 0
) -> Tuple[Any, Any, Dict[str, int]]:
    """Прочитать лист и его заголовки.

    :return: кортеж (лист, итератор строк данных, словарь ``заголовок -> позиция``)
    """
    sheets = workbook.sheetnames
    if not sheets:
        raise XlsxError(f"В файле {path.name} нет ни одного листа.")
    if sheet_index >= len(sheets):
        raise XlsxError(
            f"В файле {path.name} нет листа с индексом {sheet_index}. "
            f"Доступные листы: {', '.join(sheets)}."
        )
    sheet = workbook.worksheets[sheet_index]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise XlsxError(f"Лист «{sheet.title}» файла {path.name} пуст.")

    columns: Dict[str, int] = {}
    for position, header in enumerate(header_row):
        name = _normalize_header(header)
        if name and name not in columns:
            columns[name] = position
    return sheet, rows_iter, columns


def detect_xlsx_format(path: Path) -> str:
    """Определить формат xlsx по заголовкам первой строки.

    :param path: путь к xlsx
    :return: ``"jobs"`` (docdate/emtcodefrm), ``"manual"`` (start/end/trk)
        или ``"unknown"``
    :raises XlsxError: если файл не читается
    """
    path = Path(path)
    if not path.exists():
        raise XlsxError(f"Файл не найден: {path}")

    workbook = _open_workbook(path)
    try:
        _sheet, rows_iter, columns = _read_header(workbook, path)
        if not columns:
            return "unknown"
        if "docdate" in columns and "emtcodefrm" in columns:
            return "jobs"
        if all(name in columns for name in ("start", "end", "trk")):
            return "manual"
        return "unknown"
    finally:
        workbook.close()


def read_manual_xlsx(path: Path) -> List[ManualRow]:
    """Прочитать xlsx ручного копирования.

    Битые строки не выбрасываются: причина пишется в :attr:`ManualRow.error`,
    чтобы оператор увидел проблему в таблице.

    :param path: путь к xlsx
    :return: список :class:`ManualRow`
    :raises XlsxError: если файл не читается или в нём нет нужных колонок
    """
    path = Path(path)
    if not path.exists():
        raise XlsxError(f"Файл не найден: {path}")

    workbook = _open_workbook(path)
    try:
        sheet, rows_iter, columns = _read_header(workbook, path)
        missing = [name for name in MANUAL_COLUMNS if name not in columns]
        if missing:
            raise XlsxError(
                f"В файле {path.name} не найдены колонки ручного копирования: "
                f"{', '.join(missing)}. Ожидаются заголовки: "
                f"{', '.join(MANUAL_COLUMNS)}."
            )

        result: List[ManualRow] = []
        for row_number, values in enumerate(rows_iter, start=2):
            if values is None:
                continue
            if all(
                value is None or (isinstance(value, str) and not value.strip())
                for value in values
            ):
                continue
            result.append(_build_manual_row(row_number, values, columns, path.name))

        LOGGER.info("Прочитано строк ручного xlsx из %s: %s", path.name, len(result))
        return result
    finally:
        workbook.close()


def _cell(values: Any, columns: Dict[str, int], name: str) -> Any:
    """Достать значение ячейки по имени колонки."""
    position = columns.get(name)
    if position is None or position >= len(values):
        return None
    return values[position]


def _build_manual_row(
    row_number: int, values: Any, columns: Dict[str, int], source_name: str
) -> ManualRow:
    """Собрать :class:`ManualRow` из строки листа, отметив ошибки разбора."""
    azscode = _to_int(_cell(values, columns, "azscode"))
    start = parse_dt_text(_cell(values, columns, "start"))
    end = parse_dt_text(_cell(values, columns, "end"))
    trks, trk_error = _parse_trks(_cell(values, columns, "trk"))

    errors: List[str] = []
    if start is None:
        errors.append("не удалось прочитать start")
    if end is None:
        errors.append("не удалось прочитать end")
    if start is not None and end is not None and end <= start:
        errors.append("Конец раньше начала")
    if trk_error:
        errors.append(trk_error)

    error_text = "; ".join(errors) if errors else None
    if error_text:
        LOGGER.warning(
            "Строка %s в %s импортирована с ошибкой: %s",
            row_number,
            source_name,
            error_text,
        )

    return ManualRow(
        azscode=azscode,
        start=start if start is not None else datetime.min,
        end=end if end is not None else datetime.min,
        trks=trks,
        error=error_text,
    )


def write_manual_xlsx(path: Path, rows: List[ManualRow], default_azs: int) -> None:
    """Записать строки ручного копирования в xlsx.

    :param path: путь к создаваемому файлу
    :param rows: строки ручной таблицы
    :param default_azs: номер текущей АЗС для строк с пустым ``azscode``
    :raises XlsxError: если файл не удалось сохранить
    """
    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "manual"
    sheet.append(list(MANUAL_COLUMNS))

    for row in rows:
        azscode = row.azscode if row.azscode is not None else default_azs
        sheet.append(
            [
                azscode,
                row.start.strftime(DATETIME_FORMAT),
                row.end.strftime(DATETIME_FORMAT),
                ",".join(str(trk) for trk in row.trks),
            ]
        )

    try:
        workbook.save(str(path))
    except Exception as exc:  # openpyxl/ОС кидают разные типы исключений
        raise XlsxError(f"Не удалось сохранить файл {path.name}: {exc}") from exc
    finally:
        workbook.close()

    LOGGER.info("Записано строк ручного xlsx в %s: %s", path, len(rows))
